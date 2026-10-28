# -*- coding: utf-8 -*-
"""Confere uma peca antes do protocolo: checklist anti-emenda e lista de precedentes.

Le o .docx, aplica um checklist anti-emenda (veja o README) no que da para conferir por maquina, e separa todas as citacoes de jurisprudencia
para voce conferir na fonte. Nao julga o merito da peca e nao diz se o precedente
e bom: diz o que falta olhar.

    python checar_peca.py "caminho\\da\\peca.docx"
    python checar_peca.py "caminho\\da\\pasta"        -> confere todos os .docx da pasta
    python checar_peca.py peca.docx --json
    python checar_peca.py peca.docx --sem-leis       -> pula a conferencia dos artigos de lei

Cada "art. N ... do CPC/CC/CF/lei de Goias..." citado e buscado no texto em vigor da fonte
oficial (Planalto, Casa Civil-GO, Sefaz-GO) pelo conector juridico-br: artigo inexistente ou
revogado (inteiro, ou o paragrafo/inciso citado) barra a peca. O sentido que a peca atribui ao
artigo continua com quem revisa: o --json traz o texto vigente de cada um para comparar.

Os precedentes encontrados vao para precedentes.json (ao lado deste arquivo, ou o caminho da
variavel CHECAR_PECA_PRECEDENTES), marcados como
"a conferir" ate que voce os marque com:

    python checar_peca.py --conferido "REsp 198.196/RJ" "presuncao de culpa de quem bate atras"
"""
import json
import os
import re
import sys
from datetime import datetime

BASE_PREC = os.environ.get("CHECAR_PECA_PRECEDENTES") or os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "precedentes.json")

PLACEHOLDERS = ("[PREENCHER", "[CONFERIR", "[DECIDIR", "[NOTA", "NOTA INTERNA", "XXXXX")

CITACAO = re.compile(
    r"(?:(?:AgInt\s+n[oº°]?\s+|AgRg\s+n[oº°]?\s+)?"
    r"(?:EAREsp|AREsp|REsp|RE|HC|ADI|ADC|ADPF|RMS|MS|AI|AC|Recurso\s+Inominado|Apela[çc][ãa]o(?:\s+C[íi]vel)?)"
    r"(?:\s*n[º°o]?\.?)?\s*\d[\d.\-/]{3,}(?:\s*[/-]?\s*[A-Z]{2})?"
    r"|S[úu]mula(?:\s+Vinculante)?(?:\s*n[º°o]?\.?)?\s*\d+"
    r"|Tema(?:\s+repetitivo)?(?:\s*n[º°o]?\.?)?\s*\d+)", re.I)

ARTIGO = re.compile(r"art(?:igo)?s?\.?\s*\d+", re.I)

# --- artigos de lei citados x texto em vigor (conector juridico-br) ---
# servidor do juridico-br-mcp (https://github.com/pitersonfelizardo/juridico-br-mcp): variavel
# JURIDICO_BR_MCP ou a pasta juridico-br-mcp ao lado desta
CONECTOR = os.environ.get("JURIDICO_BR_MCP") or os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "juridico-br-mcp", "server.py")
_N = r"(?:n[º°o.]*\s*)?"
LEIS_CITADAS = [   # (padrao do nome da lei logo depois do artigo, atalho do conector)
    (r"CPC|C[óo]digo de Processo Civil|Lei\s*" + _N + r"13\.105", "cpc"),
    (r"CPP|C[óo]digo de Processo Penal|Decreto-Lei\s*" + _N + r"3\.689", "cpp"),
    (r"CC|C[óo]digo Civil|Lei\s*" + _N + r"10\.406", "cc"),
    (r"CF(?:/88)?|Constitui[çc][ãa]o(?: Federal| da Rep[úu]blica)", "cf"),
    (r"CP|C[óo]digo Penal|Decreto-Lei\s*" + _N + r"2\.848", "cp"),
    (r"CDC|C[óo]digo de Defesa do Consumidor|Lei\s*" + _N + r"8\.078", "cdc"),
    (r"CLT|Consolida[çc][ãa]o das Leis do Trabalho", "clt"),
    (r"CTN|C[óo]digo Tribut[áa]rio Nacional|Lei\s*" + _N + r"5\.172", "ctn"),
    (r"CTB|C[óo]digo de Tr[âa]nsito(?: Brasileiro)?|Lei\s*" + _N + r"9\.503", "ctb"),
    (r"ECA|Estatuto da Crian[çc]a e do Adolescente|Lei\s*" + _N + r"8\.069", "eca"),
    (r"LEP|Lei de Execu[çc][ãa]o Penal|Lei\s*" + _N + r"7\.210", "lep"),
    (r"Lei Maria da Penha|Lei\s*" + _N + r"11\.340", "maria_da_penha"),
    (r"Lei de Drogas|Lei\s*" + _N + r"11\.343", "drogas"),
    (r"Estatuto do Desarmamento|Lei\s*" + _N + r"10\.826", "desarmamento"),
    (r"Lei\s*" + _N + r"9\.099", "juizados"),
    (r"Lei do Inquilinato|Lei\s*" + _N + r"8\.245", "inquilinato"),
    (r"Lei de Alimentos|Lei\s*" + _N + r"5\.478", "alimentos"),
    (r"Estatuto da (?:Advocacia|OAB)|EAOAB|Lei\s*" + _N + r"8\.906", "estatuto_oab"),
    (r"Lei\s*" + _N + r"15\.040", "seguros"),
    (r"Lei\s*" + _N + r"14\.905", "lei_14905"),
    (r"(?:DL|Decreto-Lei)\s*" + _N + r"911", "dl_911"),
    (r"Lei\s*(?:Estadual\s*)?" + _N + r"14\.376|Regimento de Custas", "go_custas"),
    (r"Lei\s*(?:Estadual\s*)?" + _N + r"11\.651|CTE(?:-GO)?|C[óo]digo Tribut[áa]rio (?:Estadual|do Estado de Goi[áa]s)", "go_cte"),
    (r"Lei\s*(?:Estadual\s*)?" + _N + r"19\.191", "go_emolumentos"),
    (r"Lei\s*(?:Estadual\s*)?" + _N + r"21\.268|C[óo]digo de Organiza[çc][ãa]o Judici[áa]ria", "go_organizacao_judiciaria"),
    (r"Lei\s*(?:Estadual\s*)?" + _N + r"16\.898", "go_consignacoes"),
    (r"Lei\s*Estadual\s*" + _N + r"(\d{1,2}\.\d{3})", "go:"),          # outra lei de Goias pelo numero
    (r"Lei(?:\s*Complementar)?\s*" + _N + r"(\d{1,2}\.\d{3}|\d{2,4})(?:/\d{2,4})?", None),  # sem atalho
]
_LEI_ALT = "|".join(f"(?:{p})" for p, _ in LEIS_CITADAS)
CITA_ARTIGO = re.compile(
    r"\bart(?:igo)?s?\.?\s+(?P<seg>\d(?:(?!\.\s)[^;:()]){0,120}?)\s*,?\s+d[oa]s?\s+(?:referid[oa]\s+)?(?P<lei>" + _LEI_ALT + r")(?!\w|\.\d)",
    re.I)


def artigos_citados(t):
    """[(atalho ou None, nome como escrito, artigo, [partes])] sem repetir."""
    vistos, saida = set(), []
    for m in CITA_ARTIGO.finditer(t):
        nome_lei = m.group("lei")
        atalho = None
        for padrao, at in LEIS_CITADAS:
            mm = re.fullmatch(padrao, nome_lei, re.I)
            if mm:
                atalho = at + mm.group(1) if at == "go:" else at
                break
        art_atual, partes = None, {}
        for tok in re.split(r",|\be\b|\bc/c\b", m.group("seg")):
            tok = tok.strip()
            if not tok or tok.lower() in ("caput", "inc.", "inciso", "incisos"):
                continue
            tok = re.sub(r"^(?:inc(?:iso)?s?\.?\s+)", "", tok, flags=re.I)
            if tok.startswith("§") or tok.lower().startswith(("par. ún", "parágrafo único", "paragrafo unico")):
                if art_atual:
                    partes[art_atual].append("parágrafo único" if "nico" in tok.lower() else tok)
            elif re.fullmatch(r"[IVXLC]+", tok):
                if art_atual:
                    partes[art_atual].append(tok)
            else:
                a = re.match(r"(\d{1,2}\.\d{3}|\d{1,4})\s*[º°o]?(?:\s*-\s*[A-Z]\b)?", tok)
                if a:
                    art_atual = re.sub(r"\s+", "", a.group(0)).replace("o", "º")
                    partes.setdefault(art_atual, [])
        for art, ps in partes.items():
            k = (atalho or nome_lei.lower(), art, tuple(ps))
            if k not in vistos:
                vistos.add(k)
                saida.append((atalho, nome_lei, art, ps))
    return saida


def conferir_artigos(citados):
    """Manda ao conector juridico-br os artigos com atalho; devolve a lista de resultados (ou erro)."""
    import shutil
    import subprocess
    pedidos = [{"lei": at, "artigo": art, "partes": ps} for at, _, art, ps in citados if at]
    if not pedidos:
        return [], ""
    if not (shutil.which("uv") and os.path.exists(CONECTOR)):
        return [], "conector juridico-br ou uv nao encontrado"
    try:
        p = subprocess.run(["uv", "run", "--script", CONECTOR, "--conferir-artigos"],
                           input=json.dumps(pedidos), capture_output=True, text=True,
                           encoding="utf-8", timeout=300)
        return json.loads(p.stdout), ""
    except (subprocess.SubprocessError, ValueError) as e:
        return [], f"conector falhou: {str(e)[:120]}"


def texto_docx(caminho):
    import docx
    d = docx.Document(caminho)
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for lin in t.rows:
            partes += [c.text for c in lin.cells]
    return "\n".join(partes)


def arquivos_da_pasta(pasta):
    saida = []
    for raiz, dirs, arqs in os.walk(pasta):
        dirs[:] = [x for x in dirs if x not in ("_indice", "__pycache__")]
        for a in arqs:
            saida.append((a.lower(), os.path.join(raiz, a)))
    return saida


def idade_dias(caminho):
    return (datetime.now() - datetime.fromtimestamp(os.path.getmtime(caminho))).days


def carregar_precedentes():
    try:
        return json.load(open(BASE_PREC, encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def guardar_precedentes(base):
    json.dump(base, open(BASE_PREC, "w", encoding="utf-8"), ensure_ascii=False, indent=1)


def chave(c):
    import unicodedata
    c = unicodedata.normalize("NFKD", c).encode("ascii", "ignore").decode()
    c = re.sub(r"\bn[o.]?\s+(?=\d)", "", c, flags=re.I)    # "n" e "no" antes do numero nao contam
    return re.sub(r"\s+", " ", c).strip().rstrip(".,;").upper()


def conferir(caminho):
    r = {"peca": os.path.basename(caminho), "x": [], "aviso": [], "ok": [], "precedentes": []}
    t = texto_docx(caminho)
    pasta = os.path.dirname(caminho)
    arqs = arquivos_da_pasta(pasta)
    nomes = [a for a, _ in arqs]

    # 1. lacunas e notas internas
    achados = [p for p in PLACEHOLDERS if p in t.upper()]
    if achados:
        r["x"].append("Sobrou no texto: " + ", ".join(achados))
    else:
        r["ok"].append("Sem [PREENCHER], [DECIDIR] ou nota interna.")

    # 2. enderecamento, data e assinatura
    primeira = t.strip().splitlines()[0].upper() if t.strip() else ""
    if not any(k in primeira for k in ("JUIZ", "JUÍZO", "JUIZO", "EXCELENT", "VARA", "JUIZADO", "TRIBUNAL")):
        r["aviso"].append("A primeira linha nao parece um enderecamento.")
    if "OAB" not in t.upper():
        r["x"].append("Falta a assinatura com OAB.")
    if not re.search(r"\d{1,2}\s+de\s+\w+\s+de\s+20\d\d", t):
        r["aviso"].append("Nao achei a data por extenso no fim da peca.")

    # 3. qualificacao e valor da causa
    if not re.search(r"\d{3}\.\d{3}\.\d{3}-\d{2}|\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}", t):
        r["aviso"].append("Nenhum CPF ou CNPJ no texto (art. 319, II, do CPC).")
    inicial = bool(re.search(r"vem\s+propor|vem\s+ajuizar", t, re.I)) or "inicial" in os.path.basename(caminho).lower()
    if inicial:
        if re.search(r"d[áa]-se\s+[àa]\s+causa|valor\s+da\s+causa", t, re.I):
            r["ok"].append("Valor da causa declarado.")
        else:
            r["x"].append("Peca parece inicial e nao tem valor da causa (art. 291 do CPC).")
        if not re.search(r"e-?mail|endere[çc]o\s+eletr[ôo]nico", t, re.I):
            r["aviso"].append("O CPC pede o e-mail das partes; onde nao houver, dizer isso na peca.")

    # 4. documentos citados x arquivos da pasta
    docs = sorted(set(re.findall(r"DOC\s*0?(\d{1,2})", t, re.I)), key=int)
    if docs:
        faltam = [d for d in docs if not any(re.search(rf"(^|[^0-9]){int(d):02d}\s*-", n) for n in nomes)]
        if faltam:
            r["aviso"].append("DOC citado sem arquivo com o mesmo numero na pasta: " + ", ".join(faltam))
        else:
            r["ok"].append(f"Os {len(docs)} documentos citados tem arquivo correspondente.")

    # 5. procuracao assinada
    proc = [c for n, c in arqs if "procura" in n and n.endswith((".pdf", ".docx"))]
    assinada = [c for c in proc if "assinad" in os.path.basename(c).lower()]
    if inicial:
        if not proc:
            r["x"].append("Nao achei procuracao na pasta.")
        elif not assinada:
            r["x"].append("Ha procuracao, mas nenhuma marcada como assinada (art. 104 do CPC).")
        else:
            r["ok"].append("Procuracao assinada na pasta.")

    # 6. comprovante de endereco recente
    comp = [c for n, c in arqs if "comprovante" in n and "endere" in n]
    if inicial:
        if not comp:
            r["aviso"].append("Nao achei comprovante de endereco na pasta.")
        elif min(idade_dias(c) for c in comp) > 90:
            r["aviso"].append("O comprovante de endereco da pasta tem mais de 90 dias.")

    # 7. gratuidade pedida sem prova
    if re.search(r"gratuidade|hipossufici", t, re.I):
        if not any("hipossufici" in n or "declaracao" in n for n in nomes):
            r["aviso"].append("Pede gratuidade e nao achei declaracao ou prova de renda na pasta (art. 99, § 2o).")

    # 8. precedentes para conferir
    base = carregar_precedentes()
    vistos = []
    for c in CITACAO.findall(t):
        k = chave(c)
        if k in vistos:
            continue
        vistos.append(k)
        reg = base.get(k, {})
        r["precedentes"].append({"citacao": c.strip(), "estado": reg.get("estado", "a conferir"),
                                 "tese": reg.get("tese", ""), "usado_em": reg.get("usado_em", [])})
        reg.setdefault("estado", "a conferir")
        reg.setdefault("tese", "")
        usos = set(reg.get("usado_em", []))
        usos.add(os.path.basename(caminho))
        reg["usado_em"] = sorted(usos)
        base[k] = reg
    guardar_precedentes(base)
    naoconf = [p for p in r["precedentes"] if p["estado"] != "conferido"]
    if r["precedentes"]:
        (r["aviso"] if naoconf else r["ok"]).append(
            f"{len(r['precedentes'])} citacao(oes) de jurisprudencia, {len(naoconf)} sem conferencia registrada.")

    # 9. artigos de lei citados x texto em vigor na fonte oficial
    r["artigos"] = []
    if "--sem-leis" not in sys.argv:
        citados = artigos_citados(t)
        resultados, erro = conferir_artigos(citados)
        it = iter(resultados)   # mesma ordem dos citados com atalho
        for at, nome, art, ps in citados:
            cit = f"art. {art}" + "".join(f", {p}" for p in ps) + f" ({nome})"
            if not at:
                r["artigos"].append({"citacao": cit, "estado": "sem atalho"})
                continue
            res = next(it, None) if resultados else None
            if res is None:
                r["artigos"].append({"citacao": cit, "estado": "nao conferido"})
            elif res.get("erro"):
                r["artigos"].append({"citacao": cit, "estado": "nao conferido", "erro": res["erro"]})
            elif not res["achou"] or res["revogado"]:
                estado = "revogado" if res["revogado"] else "nao existe"
                r["artigos"].append({"citacao": cit, "estado": estado, "fonte": res["fonte"]})
                r["x"].append(f"Art. {art} ({nome}): {'revogado' if res['revogado'] else 'nao encontrado'} na fonte oficial ({res['fonte']}).")
            else:
                revog = [p for p in res["problemas"] if p.endswith("revogado")]
                for p in revog:
                    r["x"].append(f"Art. {art}, {p} ({nome}) na fonte oficial.")
                estado = "revogado em parte" if revog else ("conferir parte" if res["problemas"] else "em vigor")
                r["artigos"].append({"citacao": cit, "estado": estado, "problemas": res["problemas"],
                                     "texto": res["texto"], "fonte": res["fonte"]})
        if erro:
            r["aviso"].append(f"Artigos de lei nao conferidos: {erro}.")
        sem = [a for a in r["artigos"] if a["estado"] in ("sem atalho", "nao conferido", "conferir parte")]
        if sem:
            r["aviso"].append(f"{len(sem)} artigo(s) de lei para conferir a mao (lei sem atalho, fonte fora do ar ou paragrafo/inciso nao achado).")
        if r["artigos"]:
            ok = sum(1 for a in r["artigos"] if a["estado"] == "em vigor")
            r["ok"].append(f"{ok} de {len(r['artigos'])} artigo(s) de lei existem e estao em vigor na fonte oficial "
                           "(o sentido atribuido a cada um ainda e com voce).")

    r["situacao"] = "NAO PROTOCOLAR" if r["x"] else ("CONFERIR" if r["aviso"] else "PRONTA")
    return r


def imprimir(r):
    print(f"\n{r['peca']}  ->  {r['situacao']}")
    for linha in r["x"]:
        print("  X  " + linha)
    for linha in r["aviso"]:
        print("  !  " + linha)
    for linha in r["ok"]:
        print("  v  " + linha)
    if r["precedentes"]:
        print("  precedentes citados:")
        for p in r["precedentes"]:
            marca = "v" if p["estado"] == "conferido" else "?"
            tese = f" - {p['tese']}" if p["tese"] else ""
            print(f"    {marca} {p['citacao']}{tese}")
    if r.get("artigos"):
        print("  artigos de lei citados:")
        marcas = {"em vigor": "v", "revogado": "X", "nao existe": "X", "revogado em parte": "X"}
        for a in r["artigos"]:
            extra = "; ".join(a.get("problemas", [])) or a.get("erro", "")
            print(f"    {marcas.get(a['estado'], '?')} {a['citacao']} - {a['estado']}" + (f" ({extra})" if extra else ""))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    arg = sys.argv[1:]
    if not arg:
        print(__doc__)
        return 1

    if arg[0] == "--conferido":
        if len(arg) < 2:
            print('uso: checar_peca.py --conferido "REsp 198.196/RJ" "tese em uma linha"')
            return 1
        base = carregar_precedentes()
        k = chave(arg[1])
        reg = base.get(k, {"usado_em": []})
        reg["estado"] = "conferido"
        reg["conferido_em"] = datetime.now().strftime("%Y-%m-%d")
        if len(arg) > 2:
            reg["tese"] = arg[2]
        base[k] = reg
        guardar_precedentes(base)
        print(f"marcado como conferido: {arg[1]}")
        return 0

    alvo = arg[0]
    pecas = []
    if os.path.isdir(alvo):
        pecas = [os.path.join(alvo, f) for f in sorted(os.listdir(alvo))
                 if f.endswith(".docx") and not f.startswith("~$")]
    elif alvo.endswith(".docx"):
        pecas = [alvo]
    else:
        print("passe um .docx ou uma pasta")
        return 1

    saida = [conferir(p) for p in pecas]
    if "--json" in arg:
        print(json.dumps(saida, ensure_ascii=False))
    else:
        for r in saida:
            imprimir(r)
        print()
    return 1 if any(r["x"] for r in saida) else 0


if __name__ == "__main__":
    sys.exit(main())
